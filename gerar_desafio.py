#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera o Desafio Profissional preenchido (DOCX + PDF) e o texto do memorial (TXT)."""
import os
from desafio_conteudo import *

BASE = os.path.dirname(os.path.abspath(__file__))
OUT_DOCX = os.path.join(BASE, "Desafio_Profissional_Ciencias_Naturais_PREENCHIDO.docx")
OUT_PDF = os.path.join(BASE, "Desafio_Profissional_Ciencias_Naturais_PREENCHIDO.pdf")
OUT_TXT = os.path.join(BASE, "MEMORIAL_ANALITICO_para_colar_no_AVA.txt")


def memorial_texto(incluir_titulo=True):
    m = MEMORIAL
    p = []
    if incluir_titulo:
        p.append(m["titulo"])
        p.append(m["subtitulo"])
        p.append("")
    secs = [
        ("RESUMO DO QUE FOI DESCOBERTO", [m["resumo"]]),
        ("CONTEXTUALIZAÇÃO DO DESAFIO", [m["contextualizacao"]]),
        ("ANÁLISE", [m["analise"]]),
        ("PROPOSTAS DE SOLUÇÃO", m["propostas"]),
        ("CONCLUSÃO REFLEXIVA", m["conclusao"]),
        ("REFERÊNCIAS", m["referencias"]),
        ("AUTOAVALIAÇÃO", [m["autoavaliacao"]]),
    ]
    for titulo, paras in secs:
        p.append(titulo)
        for t in paras:
            p.append(t)
        p.append("")
    return "\n".join(p).strip()


def memorial_sem_titulo():
    return memorial_texto(incluir_titulo=False)


# ============================================================ DOCX
def gerar_docx():
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    doc = Document()

    # margens
    for s in doc.sections:
        s.top_margin = Cm(2.5)
        s.bottom_margin = Cm(2.5)
        s.left_margin = Cm(3)
        s.right_margin = Cm(2)

    # estilo base
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.15

    def para(text, bold=False, align="justify", size=11, space_after=6, italic=False,
             color=None, indent=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "justify":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if indent:
            p.paragraph_format.left_indent = Cm(indent)
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor(*color)
        return p

    def campo_estudante(texto_cmd):
        """Texto do template indicando o campo, em negrito, como no original."""
        para(texto_cmd, bold=True, align="left")

    # rodapé institucional em todas as páginas
    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.text = RODAPE_INSTITUCIONAL
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in fp.runs:
        r.font.size = Pt(8)
        r.font.name = "Arial"

    # ---------------- capa / orientações
    para(TITULO, bold=True, align="center", size=14, space_after=12)
    para(ORIENTACOES_TITULO, bold=True, align="left")
    for t in ORIENTACOES:
        para(t)

    # ---------------- ETAPA 1
    para("")
    para(E1_TITULO, bold=True, align="left", size=12)
    para(E1_TEXTO)

    # ---------------- ETAPA 2
    para("")
    para(E2_TITULO, bold=True, align="left", size=12)
    para(E2_TEXTO)
    para("")
    campo_estudante(E2_CMD)
    campo_estudante(E2_CAMPO)
    for t in E2_RESPOSTA:
        para(t)

    # ---------------- ETAPA 3
    para("")
    para(E3_TITULO, bold=True, align="left", size=12)
    para(E3_TEXTO)
    para("")
    campo_estudante(E3_CMD)
    for t in E3_RESPOSTA:
        para("\u2022 " + t)

    # ---------------- ETAPA 4
    para("")
    para(E4_TITULO, bold=True, align="left", size=12)
    para(E4_TEXTO)
    for t in E4_TOPICOS:
        para("\u2022 " + t)
    para("")
    campo_estudante(E4_CMD)
    for t in E4_RESPOSTA:
        para("\u2022 " + t)

    # ---------------- ETAPA 5
    para("")
    para(E5_CHAMADA, bold=True, align="center", size=12)
    para("")
    para(E5_TITULO, bold=True, align="left", size=12)
    para(E5_TEXTO1)
    para(E5_TEXTO2)
    para("Para ajudar você, segue o passo a passo do que não pode faltar no Memorial Analítico "
         "(ordem recomendada, pois cada item fará parte da composição da sua nota):", bold=False)
    for t in E5_PASSOAPASSO:
        para("\u2022 " + t)
    para("Checklist rápido antes de entregar:", bold=True, align="left")
    for t in E5_CHECKLIST:
        para("\u2022 " + t)
    para(E5_AVISO, italic=True)
    para(E5_AVISO2, italic=True)
    para("")
    campo_estudante(E5_CMD)

    # Memorial dentro do template
    m = MEMORIAL
    para(m["titulo"], bold=True, align="center", size=12)
    para(m["subtitulo"], align="center", italic=True, size=10)

    def secao(t, textos):
        para(t, bold=True, align="left", size=11, space_after=4)
        for x in textos:
            para(x)

    secao("1. Resumo do que foi descoberto", [m["resumo"]])
    secao("2. Contextualização do desafio", [m["contextualizacao"]])
    secao("3. Análise", [m["analise"]])
    secao("4. Propostas de solução", m["propostas"])
    secao("5. Conclusão reflexiva", m["conclusao"])
    secao("6. Referências", m["referencias"])
    secao("7. Autoavaliação", [m["autoavaliacao"]])

    doc.save(OUT_DOCX)
    print("DOCX gerado:", OUT_DOCX)


# ============================================================ PDF
def gerar_pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER, TA_LEFT
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib import colors

    styles = {
        "titulo": ParagraphStyle("titulo", fontName="Helvetica-Bold", fontSize=14, leading=18,
                                 alignment=TA_CENTER, spaceAfter=12),
        "sub": ParagraphStyle("sub", fontName="Helvetica-Bold", fontSize=11, leading=14,
                              alignment=TA_LEFT, spaceAfter=4, spaceBefore=8),
        "corpo": ParagraphStyle("corpo", fontName="Helvetica", fontSize=10, leading=13.2,
                                alignment=TA_JUSTIFY, spaceAfter=5),
        "item": ParagraphStyle("item", fontName="Helvetica", fontSize=10, leading=13.2,
                               alignment=TA_JUSTIFY, spaceAfter=4, leftIndent=14,
                               bulletIndent=4),
        "bold": ParagraphStyle("bold", fontName="Helvetica-Bold", fontSize=10.5, leading=14,
                               alignment=TA_LEFT, spaceAfter=4),
        "italic": ParagraphStyle("italic", fontName="Helvetica-Oblique", fontSize=9.5, leading=13,
                                 alignment=TA_LEFT, spaceAfter=4),
        "center9": ParagraphStyle("center9", fontName="Helvetica", fontSize=8.5, leading=11,
                                  alignment=TA_CENTER, spaceAfter=8, textColor=colors.grey),
        "center10": ParagraphStyle("center10", fontName="Helvetica-Bold", fontSize=11.5,
                                   leading=15, alignment=TA_CENTER, spaceAfter=6),
    }

    def P(t, s="corpo"):
        return Paragraph(t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"), styles[s])

    def esc(t):
        return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    story = []
    story.append(P(TITULO, "titulo"))
    story.append(P(ORIENTACOES_TITULO, "bold"))
    for t in ORIENTACOES:
        story.append(Paragraph(esc(t), styles["corpo"]))

    story.append(P(E1_TITULO, "sub"))
    story.append(Paragraph(esc(E1_TEXTO), styles["corpo"]))

    story.append(P(E2_TITULO, "sub"))
    story.append(Paragraph(esc(E2_TEXTO), styles["corpo"]))
    story.append(Paragraph(esc(E2_CMD), styles["bold"]))
    story.append(Paragraph(esc(E2_CAMPO), styles["bold"]))
    for t in E2_RESPOSTA:
        story.append(Paragraph(esc(t), styles["corpo"]))

    story.append(P(E3_TITULO, "sub"))
    story.append(Paragraph(esc(E3_TEXTO), styles["corpo"]))
    story.append(Paragraph(esc(E3_CMD), styles["bold"]))
    for t in E3_RESPOSTA:
        story.append(Paragraph("&bull; " + esc(t), styles["item"]))

    story.append(P(E4_TITULO, "sub"))
    story.append(Paragraph(esc(E4_TEXTO), styles["corpo"]))
    for t in E4_TOPICOS:
        story.append(Paragraph("&bull; " + esc(t), styles["item"]))
    story.append(Paragraph(esc(E4_CMD), styles["bold"]))
    for t in E4_RESPOSTA:
        story.append(Paragraph("&bull; " + esc(t), styles["item"]))

    story.append(Paragraph(esc(E5_CHAMADA), styles["center10"]))
    story.append(P(E5_TITULO, "sub"))
    story.append(Paragraph(esc(E5_TEXTO1), styles["corpo"]))
    story.append(Paragraph(esc(E5_TEXTO2), styles["corpo"]))
    story.append(Paragraph("Para ajudar você, segue o passo a passo do que não pode faltar no Memorial "
                           "Analítico (ordem recomendada, pois cada item fará parte da composição da "
                           "sua nota):", styles["corpo"]))
    for t in E5_PASSOAPASSO:
        story.append(Paragraph("&bull; " + esc(t), styles["item"]))
    story.append(Paragraph("Checklist rápido antes de entregar:", styles["bold"]))
    for t in E5_CHECKLIST:
        story.append(Paragraph("&bull; " + esc(t), styles["item"]))
    story.append(Paragraph(esc(E5_AVISO), styles["italic"]))
    story.append(Paragraph(esc(E5_AVISO2), styles["italic"]))
    story.append(Paragraph(esc(E5_CMD), styles["bold"]))

    m = MEMORIAL
    story.append(Paragraph(esc(m["titulo"]), styles["center10"]))
    story.append(Paragraph(esc(m["subtitulo"]), styles["italic"]))

    def secao(t, textos):
        story.append(P(t, "sub"))
        for x in textos:
            story.append(Paragraph(esc(x), styles["corpo"]))

    secao("1. Resumo do que foi descoberto", [m["resumo"]])
    secao("2. Contextualização do desafio", [m["contextualizacao"]])
    secao("3. Análise", [m["analise"]])
    secao("4. Propostas de solução", m["propostas"])
    secao("5. Conclusão reflexiva", m["conclusao"])
    secao("6. Referências", m["referencias"])
    secao("7. Autoavaliação", [m["autoavaliacao"]])

    from reportlab.lib import colors as _colors

    def _rodape(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 6.8)
        canvas.setFillColor(_colors.grey)
        canvas.drawCentredString(A4[0] / 2.0, 1.25 * cm, RODAPE_INSTITUCIONAL)
        canvas.restoreState()

    doc = SimpleDocTemplate(OUT_PDF, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm,
                            leftMargin=2.5 * cm, rightMargin=2 * cm,
                            title="Desafio Profissional - Ciencias Naturais e Saude Infantil",
                            author="Estudante")
    doc.build(story, onFirstPage=_rodape, onLaterPages=_rodape)
    print("PDF gerado:", OUT_PDF)


# ============================================================ TXT
def gerar_txt():
    texto = memorial_texto(incluir_titulo=False)
    with open(OUT_TXT, "w", encoding="utf-8") as f:
        f.write(texto)
    print("TXT gerado:", OUT_TXT)
    print("\n=== CONTAGEM DE CARACTERES DO MEMORIAL (inclui seções e referências) ===")
    print("Sem título: ", len(texto))
    com = memorial_texto(incluir_titulo=True)
    print("Com título: ", len(com))
    return texto


if __name__ == "__main__":
    gerar_docx()
    gerar_pdf()
    gerar_txt()
